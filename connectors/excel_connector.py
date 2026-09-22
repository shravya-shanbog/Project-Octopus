# ==========================================================
# PROJECT OCTOPUS — EXCEL CONNECTOR
# Pandas / NumPy free implementation
# ==========================================================

from pathlib import Path
import hashlib

from openpyxl import load_workbook

from connectors.base_connector import BaseConnector


class ExcelConnector(BaseConnector):
    """
    Connector for Excel workbooks.

    Uses OpenPyXL instead of Pandas.
    Supports .xlsx files.
    """

    def __init__(self, config: dict):

        super().__init__(config)

        self.file_path = Path(
            config.get("file_path", "")
        )

        self.sheet_name = config.get(
            "sheet_name"
        )

        self.rows = []

        self.connected = False

    # ------------------------------------------------------
    # CONNECT
    # ------------------------------------------------------

    def connect(self) -> None:

        if not self.file_path:

            raise ValueError(
                "Excel file path is required."
            )

        if not self.file_path.exists():

            raise FileNotFoundError(
                f"Excel file not found: {self.file_path}"
            )

        if self.file_path.suffix.lower() not in {
            ".xlsx",
            ".xlsm",
        }:

            raise ValueError(
                "Unsupported Excel format. "
                "Use .xlsx or .xlsm."
            )

        self.connected = True

    # ------------------------------------------------------
    # TEST CONNECTION
    # ------------------------------------------------------

    def test_connection(self) -> bool:

        try:

            self.connect()

            workbook = load_workbook(
                filename=self.file_path,
                read_only=True,
                data_only=True,
            )

            workbook.close()

            return True

        except Exception:

            return False

    # ------------------------------------------------------
    # GET WORKSHEET
    # ------------------------------------------------------

    def _get_sheet(
        self,
        workbook,
    ):

        if self.sheet_name:

            if self.sheet_name not in workbook.sheetnames:

                raise ValueError(
                    f"Worksheet '{self.sheet_name}' "
                    "was not found."
                )

            return workbook[
                self.sheet_name
            ]

        if not workbook.sheetnames:

            raise ValueError(
                "Excel workbook contains no worksheets."
            )

        return workbook[
            workbook.sheetnames[0]
        ]

    # ------------------------------------------------------
    # DISCOVER
    # ------------------------------------------------------

    def discover(self) -> dict:

        if not self.connected:
            self.connect()

        workbook = load_workbook(
            filename=self.file_path,
            read_only=True,
            data_only=True,
        )

        try:

            sheet = self._get_sheet(
                workbook
            )

            rows = sheet.iter_rows(
                values_only=True
            )

            header_row = next(
                rows,
                None,
            )

            if not header_row:

                return {
                    "file_name": (
                        self.file_path.name
                    ),
                    "file_path": str(
                        self.file_path
                    ),
                    "source_type": "Excel",
                    "sheet_name": (
                        sheet.title
                    ),
                    "columns": [],
                    "column_count": 0,
                    "row_count": 0,
                }

            columns = []

            for index, value in enumerate(
                header_row
            ):

                if value is None:

                    column_name = (
                        f"column_{index + 1}"
                    )

                else:

                    column_name = str(
                        value
                    ).strip()

                columns.append(
                    column_name
                )

            row_count = 0

            for row in rows:

                if any(
                    value not in (
                        None,
                        "",
                    )
                    for value in row
                ):

                    row_count += 1

            return {
                "file_name": (
                    self.file_path.name
                ),
                "file_path": str(
                    self.file_path
                ),
                "source_type": "Excel",
                "sheet_name": (
                    sheet.title
                ),
                "columns": columns,
                "column_count": len(
                    columns
                ),
                "row_count": row_count,
            }

        finally:

            workbook.close()

    # ------------------------------------------------------
    # EXTRACT
    # ------------------------------------------------------

    def extract(
        self,
        target=None,
    ) -> list[dict]:

        if not self.connected:
            self.connect()

        requested_sheet = (
            target
            or self.sheet_name
        )

        workbook = load_workbook(
            filename=self.file_path,
            read_only=True,
            data_only=True,
        )

        try:

            if requested_sheet:

                if requested_sheet not in workbook.sheetnames:

                    raise ValueError(
                        f"Worksheet '{requested_sheet}' "
                        "was not found."
                    )

                sheet = workbook[
                    requested_sheet
                ]

            else:

                if not workbook.sheetnames:

                    raise ValueError(
                        "Excel workbook contains no worksheets."
                    )

                sheet = workbook[
                    workbook.sheetnames[0]
                ]

            iterator = sheet.iter_rows(
                values_only=True
            )

            header_row = next(
                iterator,
                None,
            )

            if not header_row:

                raise ValueError(
                    "Excel worksheet does not "
                    "contain a header row."
                )

            headers = []

            for index, value in enumerate(
                header_row
            ):

                if value is None:

                    header = (
                        f"column_{index + 1}"
                    )

                else:

                    header = str(
                        value
                    ).strip()

                headers.append(
                    header
                )

            rows = []

            for raw_row in iterator:

                row = {}

                for index, header in enumerate(
                    headers
                ):

                    if index < len(
                        raw_row
                    ):

                        value = raw_row[
                            index
                        ]

                    else:

                        value = None

                    # Preserve actual datetime
                    # objects from OpenPyXL.
                    #
                    # Other values are converted
                    # to clean strings.

                    if (
                        value is not None
                        and not hasattr(
                            value,
                            "isoformat",
                        )
                    ):

                        value = str(
                            value
                        ).strip()

                    row[header] = value

                if any(
                    value not in (
                        None,
                        "",
                    )
                    for value in row.values()
                ):

                    rows.append(
                        row
                    )

            self.rows = rows

            return rows

        finally:

            workbook.close()

    # ------------------------------------------------------
    # SOURCE INFORMATION
    # ------------------------------------------------------

    def get_source_info(self) -> dict:

        path_value = str(
            self.file_path.resolve()
        )

        source_hash = hashlib.sha256(
            path_value.encode(
                "utf-8"
            )
        ).hexdigest()[:16]

        info = super().get_source_info()

        info.update(
            {
                "source_id": (
                    f"SRC-{source_hash}"
                ),
                "source_type": "Excel",
                "source_name": (
                    self.file_path.name
                ),
                "file_name": (
                    self.file_path.name
                ),
                "file_path": path_value,
                "connection_reference": (
                    path_value
                ),
                "sheet_name": (
                    self.sheet_name
                ),
            }
        )

        return info

    # ------------------------------------------------------
    # CLOSE
    # ------------------------------------------------------

    def close(self) -> None:

        self.rows = []

        self.connected = False